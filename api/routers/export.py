from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from api.database import get_db_connection
from api.security import verify_token
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from decimal import Decimal

router = APIRouter()

@router.get("/export-excel/", dependencies=[Depends(verify_token)])
def export_monthly_report(month: str = Query(..., description="Format: YYYY-MM")):
    conn = get_db_connection()
    cursor = conn.cursor()
    
    try:
        # 1. Fetch Financial Settings
        cursor.execute("SELECT vat_rate, mgmt_fee_pct, arrears_fee, shortcode_fee FROM company LIMIT 1")
        settings = cursor.fetchone()
        vat_rate = float(settings[0]) / 100
        mgmt_fee_pct = float(settings[1]) / 100
        arrears_fee = float(settings[2])
        shortcode_fee = float(settings[3])

        # 2. Fetch Data for the specified month
        cursor.execute("""
            SELECT t.id, t.first_name, t.last_name, t.property_name, 
                   tx.transaction_type, tx.amount, tx.reference, tx.created_at
            FROM transactions tx
            JOIN wallets w ON tx.wallet_id = w.id
            JOIN tenants t ON w.tenant_id = t.id
            WHERE TO_CHAR(tx.created_at, 'YYYY-MM') = %s
            ORDER BY tx.created_at ASC
        """, (month,))
        transactions = cursor.fetchall()

        # 3. Process Data
        energy_data = []
        water_data = []
        payments_data = []
        total_elec_cost = 0.0
        total_water_cost = 0.0
        total_payments = 0.0
        total_arrears_loaded = 0

        for txn in transactions:
            txn_type = txn[4]
            amount = float(txn[5])
            ref = txn[6] or ""
            date = txn[7].strftime("%Y-%m-%d")
            
            # It's a Bill (Usage)
            if txn_type in ('ELECTRICITY_BILL', 'ELECTRICITY_USAGE'):
                net = amount / (1 + vat_rate)
                vat = amount - net
                energy_data.append([txn[0], f"{txn[1]} {txn[2]}", txn[3], date, "Utilities", net, vat, amount, ref])
                total_elec_cost += amount
            elif txn_type in ('WATER_BILL', 'WATER_USAGE'):
                net = amount / (1 + vat_rate)
                vat = amount - net
                water_data.append([txn[0], f"{txn[1]} {txn[2]}", txn[3], date, "Utilities", net, vat, amount, ref])
                total_water_cost += amount
            # It's a Payment
            elif txn_type in ('WALLET_TOPUP', 'RENT_PAYMENT', 'UTILITY_PAYMENT'):
                payments_data.append([date, txn[0], f"{txn[1]} {txn[2]}", txn[3], "Payment", amount, 0, amount])
                total_payments += amount
            # It's an Arrears Load
            elif txn_type == 'ADJUSTMENT_RENT':
                total_arrears_loaded += 1

        # 4. Calculate Management Fees
        mgmt_fee_amount = total_payments * mgmt_fee_pct
        arrears_fee_amount = total_arrears_loaded * arrears_fee
        subtotal = mgmt_fee_amount + arrears_fee_amount + shortcode_fee
        vat_on_fees = subtotal * vat_rate
        total_invoice = subtotal + vat_on_fees

        # 5. Create Excel Workbook
        wb = Workbook()
        
        # --- Sheet 1: Summary ---
        ws1 = wb.active
        ws1.title = "Summary"
        ws1.append(["Monthly Report - " + month])
        ws1.append([""])
        ws1.append(["Consumption", "Units", "Total Cost (Incl VAT)"])
        ws1.append(["Electricity", "", total_elec_cost])
        ws1.append(["Water", "", total_water_cost])
        ws1.append([""])
        ws1.append(["Total Consumed & Recovered", "", total_elec_cost + total_water_cost])
        ws1.append([""])
        ws1.append(["Management Fee Invoice"])
        ws1.append([f"Transaction Fees ({mgmt_fee_pct*100}% of {total_payments})", mgmt_fee_amount])
        ws1.append([f"Arrears/Rent Recovery Fee ({total_arrears_loaded} loads @ R{arrears_fee})", arrears_fee_amount])
        ws1.append(["Monthly Short Code Enquiry Fee", shortcode_fee])
        ws1.append(["Subtotal", subtotal])
        ws1.append([f"VAT ({vat_rate*100}%)", vat_on_fees])
        ws1.append(["Total Due", total_invoice])

        # --- Sheet 2: Energy ---
        ws2 = wb.create_sheet("Energy")
        ws2.append(["Tenant ID", "Tenant Name", "Property", "Date", "Type", "Net", "VAT", "Total", "Reference"])
        for row in energy_data:
            ws2.append(row)

        # --- Sheet 3: Water ---
        ws3 = wb.create_sheet("Water")
        ws3.append(["Tenant ID", "Tenant Name", "Property", "Date", "Type", "Net", "VAT", "Total", "Reference"])
        for row in water_data:
            ws3.append(row)

        # --- Sheet 4: Payments ---
        ws4 = wb.create_sheet("Payments")
        ws4.append(["Date", "Tenant ID", "Tenant Name", "Property", "Action", "Net", "VAT", "Total"])
        for row in payments_data:
            ws4.append(row)

        # 6. Save to BytesIO (Memory) for download
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)

        headers = {
            'Content-Disposition': f'attachment; filename="Monthly_Report_{month}.xlsx"'
        }
        return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

    except Exception as e:
        return {"error": str(e)}
    finally:
        cursor.close()
        conn.close()